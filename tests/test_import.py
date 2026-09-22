from docx import Document

from app.services.docx_import import import_prologue


def test_import_preserves_paragraphs_and_emphasis(tmp_path):
    path = tmp_path / 'story.docx'
    d = Document()
    d.add_paragraph('«Моя слабость, моя боль»')
    d.add_paragraph('роман')
    d.add_paragraph('Посвящение.').runs[0].italic = True
    d.add_paragraph('Пролог')
    p = d.add_paragraph('Первое предложение. ')
    p.add_run('Выделенная мысль.').italic = True
    d.add_paragraph('— Реплика персонажа.')
    d.save(path)
    result = import_prologue(path)
    assert result['title'] == 'Моя слабость, моя боль'
    assert result['source_paragraphs'] == 6 and result['prose_paragraphs'] == 2
    assert '<blockquote><p><em>Посвящение.</em></p></blockquote>' in result['body']
    assert '<em>Выделенная мысль.</em>' in result['body']
    assert '<p>— Реплика персонажа.</p>' in result['body']
